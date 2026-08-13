from __future__ import annotations

import json
import os
import re
import subprocess
from collections import defaultdict
from datetime import datetime
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "BOP_Realty_HRMS_Complete_Technical_and_Functional_Documentation.docx"
TODAY = "10 August 2026"
VERSION = "1.2.0"
BUILD = "20260810.1"
SCHEMA = "20260810000000"

PURPLE = "6D1B6D"
PURPLE_DARK = "4A0F4A"
PURPLE_PALE = "F0EAF5"
INK = "1A1028"
GREY = "6B6478"
LIGHT_GREY = "EDE5F2"
WHITE = "FFFFFF"
GREEN = "12966E"
AMBER = "D97706"
RED = "C8252D"
BLUE = "1A5FD4"


def read(path: str | Path) -> str:
    return (ROOT / path).read_text(encoding="utf-8", errors="replace")


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell, **kwargs) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        if edge not in kwargs:
            continue
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        for key in ("val", "sz", "space", "color"):
            if key in kwargs[edge]:
                element.set(qn("w:" + key), str(kwargs[edge][key]))


def set_cell_text(cell, text: object, bold: bool = False, color: str | None = None, size: float = 8.5) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.space_before = Pt(0)
    run = p.add_run(str(text if text is not None else ""))
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Aptos"
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc: Document, headers: list[str], rows: Iterable[Iterable[object]], widths: list[float] | None = None,
              font_size: float = 8.25, repeat_header: bool = True):
    rows = list(rows)
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, PURPLE_DARK)
        set_cell_text(cell, header, bold=True, color=WHITE, size=font_size)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths:
            cell.width = Inches(widths[i])
    if repeat_header:
        tr_pr = table.rows[0]._tr.get_or_add_trPr()
        tbl_header = OxmlElement("w:tblHeader")
        tbl_header.set(qn("w:val"), "true")
        tr_pr.append(tbl_header)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=font_size)
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            if ridx % 2:
                set_cell_shading(cells[i], "FAF8FC")
            if widths:
                cells[i].width = Inches(widths[i])
            set_cell_border(cells[i], bottom={"val": "single", "sz": "3", "color": LIGHT_GREY})
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_bullets(doc: Document, items: Iterable[str], level: int = 0) -> None:
    for item in items:
        p = doc.add_paragraph(style="List Bullet" if level == 0 else "List Bullet 2")
        p.add_run(item)


def add_numbered(doc: Document, items: Iterable[str]) -> None:
    for item in items:
        doc.add_paragraph(item, style="List Number")


def add_note(doc: Document, title: str, body: str, color: str = BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    set_cell_shading(cell, "F4F7FD" if color == BLUE else "FFF7ED")
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(color)
    p.add_run(body)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_code(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="Code Block")
    p.add_run(text)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_toc(doc: Document) -> None:
    p = doc.add_paragraph()
    run = p.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "Right-click and select Update Field to refresh this table of contents."
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_sep, placeholder, fld_end])


def configure_document(doc: Document) -> None:
    sec = doc.sections[0]
    sec.top_margin = Inches(0.65)
    sec.bottom_margin = Inches(0.65)
    sec.left_margin = Inches(0.72)
    sec.right_margin = Inches(0.72)
    sec.header_distance = Inches(0.25)
    sec.footer_distance = Inches(0.25)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.08
    for name, size, color in (("Title", 34, PURPLE_DARK), ("Heading 1", 22, PURPLE_DARK),
                              ("Heading 2", 15, PURPLE), ("Heading 3", 11.5, INK)):
        s = styles[name]
        s.font.name = "Aptos Display"
        s.font.size = Pt(size)
        s.font.bold = True
        s.font.color.rgb = RGBColor.from_string(color)
        s.paragraph_format.keep_with_next = True
        s.paragraph_format.space_before = Pt(10)
        s.paragraph_format.space_after = Pt(5)
    if "Code Block" not in styles:
        code = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(7.8)
        code.font.color.rgb = RGBColor.from_string("2E0630")
        code.paragraph_format.left_indent = Inches(0.18)
        code.paragraph_format.right_indent = Inches(0.12)
        code.paragraph_format.space_after = Pt(6)
        p_pr = code.element.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F7F2F8")
        p_pr.append(shd)
    header = sec.header.paragraphs[0]
    header.text = "BOP REALTY  |  HRMS SYSTEM DOCUMENTATION  |  v1.2.0"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(7.5)
    header.runs[0].font.color.rgb = RGBColor.from_string(GREY)
    footer = sec.footer.paragraphs[0]
    add_page_number(footer)
    for r in footer.runs:
        r.font.size = Pt(7.5)
        r.font.color.rgb = RGBColor.from_string(GREY)
    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)


def clean_identifier(value: str) -> str:
    return value.replace('"', "").split(".")[-1].strip()


def sql_files() -> list[Path]:
    return sorted((ROOT / "supabase" / "migrations").glob("*.sql")) + sorted((ROOT / "scripts").glob("*.sql"))


def collect_sql_objects():
    objects: dict[str, dict[str, set[str]]] = {
        "table": defaultdict(set), "view": defaultdict(set), "function": defaultdict(set), "policy": defaultdict(set)
    }
    for path in sql_files():
        text = path.read_text(encoding="utf-8", errors="replace")
        rel = path.relative_to(ROOT).as_posix()
        for match in re.finditer(r"create\s+table\s+(?:if\s+not\s+exists\s+)?((?:public\.)?[\"\w]+)", text, re.I):
            name = clean_identifier(match.group(1))
            if name not in {"public", "to"}:
                objects["table"][name].add(rel)
        for match in re.finditer(r"create\s+(?:or\s+replace\s+)?(?:materialized\s+)?view\s+((?:public\.)?[\"\w]+)", text, re.I):
            objects["view"][clean_identifier(match.group(1))].add(rel)
        for match in re.finditer(r"create\s+(?:or\s+replace\s+)?function\s+((?:public\.)?[\"\w]+)", text, re.I):
            name = clean_identifier(match.group(1))
            if name != "to":
                objects["function"][name].add(rel)
        for match in re.finditer(r"create\s+policy\s+(?:\"([^\"]+)\"|([\w]+)).*?\bon\s+((?:public\.)?[\"\w]+|storage\.objects)", text, re.I | re.S):
            policy = match.group(1) or match.group(2)
            table = clean_identifier(match.group(3))
            objects["policy"][f"{table}: {policy}"].add(rel)
    return objects


def ts_refs(paths: Iterable[Path]):
    tables, rpcs, functions = defaultdict(set), defaultdict(set), defaultdict(set)
    for path in paths:
        text = path.read_text(encoding="utf-8", errors="replace")
        rel = path.relative_to(ROOT).as_posix()
        for name in re.findall(r"\.from\(\s*[\"']([^\"']+)", text):
            tables[name].add(rel)
        for name in re.findall(r"\.rpc\(\s*[\"']([^\"']+)", text):
            rpcs[name].add(rel)
        for name in re.findall(r"functions\.invoke\(\s*[\"']([^\"']+)", text):
            functions[name].add(rel)
    return tables, rpcs, functions


def collect_screens():
    result = []
    active_names = set()
    for nav in (ROOT / "navigation").glob("*.tsx"):
        text = nav.read_text(encoding="utf-8", errors="replace")
        active_names.update(re.findall(r"@/screens/(?:admin/)?([A-Za-z0-9_-]+)", text))
    active_names.update(["LoginScreen", "DashboardScreen"])
    descriptions = {
        "LoginScreen": "Phone-number OTP request and verification; establishes the Supabase session.",
        "DashboardScreen": "Employee landing page, attendance state/actions, summary cards, events, quick links and tracking health.",
        "AttendanceScreen": "Monthly attendance detail, official evidence, checkout, fraud diagnostics and regularization submission.",
        "FieldVisitScreen": "Create, start, arrive/check in and check out of office/site/meeting journeys; fallback evidence capture.",
        "TasksScreen": "Employee event/task list, task detail, status updates and event attendance actions.",
        "LeaveScreen": "Leave balance, calendar, application history, submission and cancellation.",
        "ProfileScreen": "Profile/avatar maintenance, account links, device sessions, theme and logout.",
        "MyTeamScreen": "Manager team status, attendance/tasks/visits and latest-location visualization.",
        "NotificationsScreen": "In-app notifications, event acknowledgement, read state and deep-link routing.",
        "ActivityLoggingScreen": "Structured business activity entry with office/site context and quality analysis.",
        "MyActivityLogsScreen": "Employee activity history and navigation to new activity entry.",
        "RaiseTicketScreen": "Support ticket creation and personal ticket history.",
        "LocationTrackingSetupScreen": "Permission education, tracking health, settings and restoration diagnostics.",
        "UpcomingBirthdaysScreen": "Thirty-day employee birthday list from a database view.",
        "AdminDashboardScreen": "HR Panel summary and launch point for administrative modules.",
        "UserManagementScreen": "Protected employee create/edit/import/activate/deactivate management.",
        "TeamHierarchyScreen": "Reporting hierarchy editing, Sales Manager positions/direct reports and bulk hierarchy import.",
        "AttendanceControlScreenNew": "Canonical daily attendance control, evidence comparison, warnings, correction and exports.",
        "AdvancedAttendanceReportScreen": "Detailed attendance reporting, filters, summaries and download centre.",
        "AttendanceCalendarScreen": "Holiday and weekly-off configuration scoped by company/office/team/employee.",
        "AttendanceRegularizationScreen": "HR review of attendance correction requests.",
        "CheckInLogsScreen": "App check-in attempts and committed attendance evidence inspection.",
        "BiometricUserMatchScreen": "Biometric identity mapping, unmatched punches and reprocessing.",
        "SmartOfficeSecureSyncScreen": "Protected biometric synchronization controls and status.",
        "SmartOfficeEmployeeImport": "Spreadsheet employee import, matching, batch assignment and validation.",
        "SmartOfficeImportedData": "Imported SmartOffice punch/profile inspection.",
        "LeaveManagementScreen": "Leave review, calendar detail and HR comp-off credit.",
        "LeaveBalanceScreen": "Ledger-backed balances and idempotent HR adjustments.",
        "FieldVisitAdminScreen": "Organization field-visit report and evidence review.",
        "LiveTrackingScreen": "Authorized live location map, alerts, ping requests and tracking diagnostics.",
        "SuspiciousAttendanceScreen": "Fraud-event queue and administrative review.",
        "EventManagementScreen": "Targeted event creation/deletion, audience selection, attendance and photos.",
        "TaskProjectControlScreen": "Organization task creation, assignment, status and project-style control.",
        "NotificationManagementScreen": "Administrative in-app notification composition and recipient selection.",
        "TicketManagementScreen": "Support ticket triage, resolution and closure.",
        "OfficeSettingsScreen": "Office geofence/configuration management and configuration issue checks.",
        "SiteSettingsScreen": "Predefined destination/site geofence management.",
        "AuditLogsScreen": "Administrative audit record inspection.",
        "ActivityAnalysisScreen": "Cross-employee activity reporting and export.",
        "TVDisplayScreen": "Large-screen employee attendance/occupancy/tracking display with realtime refresh.",
        "ReportsHubScreen": "Role-protected report directory.",
    }
    for path in sorted((ROOT / "screens").rglob("*.tsx")):
        text = path.read_text(encoding="utf-8", errors="replace")
        name = path.stem
        tables = sorted(set(re.findall(r"\.from\(\s*[\"']([^\"']+)", text)))
        rpcs = sorted(set(re.findall(r"\.rpc\(\s*[\"']([^\"']+)", text)))
        imports = sorted(set(re.findall(r"@/services/([^\"']+)", text)))
        active = name in active_names
        if name in {"ReportsHubScreen", "SmartOfficeSecureSyncScreen", "SmartOfficeEmployeeImport", "SmartOfficeImportedData"}:
            active = True
        result.append({"path": path.relative_to(ROOT).as_posix(), "name": name, "active": active,
                       "lines": text.count("\n") + 1, "tables": tables, "rpcs": rpcs, "services": imports,
                       "description": descriptions.get(name, "Legacy, placeholder, alternate, or specialized screen; inspect route status before use.")})
    return result


def collect_services():
    descriptions = {
        "attendance.service.ts": "Office/geofence lookup and protected check-in, reconciled checkout, undo, correction, policy and warnings.",
        "canonical-attendance.service.ts": "Canonical V3 daily attendance query and UI compatibility mapping.",
        "field-visit.service.ts": "Own-visit lifecycle and selected-site RPC mutations.",
        "current-location-visit.service.ts": "Photo/location integrity evidence, encrypted native retry queue and reviewer operations.",
        "leave.service.ts": "Leave balances, submission/cancellation/review and scoped calendar lookup.",
        "event.service.ts": "Targeted events, announcements, attendance check-in/out and attendee listing.",
        "activity.service.ts": "Activity CRUD, daily employee logs and analytics.",
        "activity-analysis.service.ts": "Client-side entity extraction, categorization and low-quality detection.",
        "dashboard-summary.service.ts": "Cached employee/admin summary RPC facade.",
        "smartoffice.service.ts": "SmartOffice profile/device-log synchronization, mapping and bulk import.",
        "attendance-import.service.ts": "XLS/XLSX parsing, normalization, employee matching and secure batch import.",
        "biometric-sync-secure.service.ts": "Protected Edge Function sync with progress polling and durable web state.",
        "device-session.service.ts": "Device registration, list, revoke and close RPCs.",
        "support-ticket.service.ts": "Ticket create/list/status workflow; messaging hooks currently disabled.",
        "business-clock.service.ts": "Cached authoritative Asia/Kolkata server clock.",
        "tracking-admin.service.ts": "Admin fraud review, alert resolution and explicit ping requests.",
        "predefined-site.service.ts": "Site list and client-side Haversine/geofence selection.",
        "office.service.ts": "Office directory query.",
        "task.service.ts": "Small pending-task query helper.",
        "notification.service.ts": "Unread notification count helper.",
        "whatsapp.service.ts": "Disabled placeholder; refuses client-side provider delivery.",
        "sms.service.ts": "Disabled placeholder; OTP delivery must use the server function.",
    }
    result = []
    for path in sorted((ROOT / "services").rglob("*.ts")):
        if "__tests__" in path.parts:
            continue
        text = path.read_text(encoding="utf-8", errors="replace")
        tables = sorted(set(re.findall(r"\.from\(\s*[\"']([^\"']+)", text)))
        rpcs = sorted(set(re.findall(r"\.rpc\(\s*[\"']([^\"']+)", text)))
        edge = sorted(set(re.findall(r"functions\.invoke\(\s*[\"']([^\"']+)", text)))
        result.append({"path": path.relative_to(ROOT).as_posix(), "name": path.name,
                       "lines": text.count("\n") + 1, "tables": tables, "rpcs": rpcs, "edge": edge,
                       "description": descriptions.get(path.name, "Supporting service or platform-specific implementation; see data dependencies.")})
    return result


TABLE_DOMAINS = {
    "profiles": "Identity / HR master", "roles": "Authorization", "offices": "Configuration / geofence",
    "attendance_logs": "Attendance", "attendance_policies": "Attendance configuration", "attendance_regularization_requests": "Attendance review",
    "attendance_fraud_events": "Fraud", "attendance_audit_logs": "Audit", "checkin_attempts": "Attendance telemetry",
    "field_visits": "Field operations", "field_visit_location_audit": "Field evidence", "location_tracking": "Location tracking",
    "tracking_sessions": "Location tracking", "journey_events": "Location tracking", "tracking_health_events": "Location diagnostics",
    "leave_requests": "Leave", "leave_balance": "Leave", "leave_transactions": "Leave ledger", "comp_off_earnings": "Leave / comp-off",
    "tasks": "Work management", "activity_logs": "Work activity", "support_tickets": "Support", "notifications": "Messaging",
    "events": "Events", "event_checkins": "Events", "event_participants": "Events / audience",
    "smartoffice_device_logs": "Biometric", "employee_biometric_mappings": "Biometric", "biometric_device_punches": "Biometric",
    "device_push_tokens": "Push messaging", "notification_push_outbox": "Push messaging", "user_device_sessions": "Session security",
    "audit_logs": "Audit", "hrms_audit_events": "Audit", "client_operational_errors": "Observability",
}


CORE_COLUMNS = {
    "profiles": "id, name, email, phone, role, employee_number, avatar, is_active, office_id, reports_to, team/team_id, department, show_on_tv, attendance_reporting_enabled, is_service_account, created_at",
    "attendance_logs": "id, user_id, date, checkin_time, checkout_time, status, checkin/checkout locations, office IDs, source fields, biometric evidence, warning fields, correction/audit fields",
    "offices": "id, name, address, latitude, longitude, attendance_radius/allowed_radius_meters, default_attendance_time, is_active, capacity",
    "tasks": "id, title, description, assigned_to/assigned_by, status, priority, due_date, event_id, timestamps",
    "notifications": "id, user_id, title, message/body, type, read/is_read, related entity/deep-link data, created_at",
    "leave_balance": "id, user_id, sick, casual, earned, comp_off, updated_at",
    "leave_requests": "id, user_id, leave_type, start_date, end_date, days, reason, status, reviewer fields, timestamps",
    "field_visits": "id, user_id, type/status, destination/site/office coordinates, scheduled/visit date, check-in/out timestamps and coordinates, journey origin, photo evidence and integrity metadata",
    "location_tracking": "id, user_id, attendance/visit/session references, latitude, longitude, accuracy, speed/heading, status, recorded_at, client event ID, distance/context",
    "events": "id, name/title, description, venue and coordinates, radius, start/end/duty times, active/published flags, audience rules and timestamps",
    "event_checkins": "event/user IDs, check-in/out timestamps and coordinates, distance/accuracy, source and attendance linkage",
    "activity_logs": "id, user_id, activity_type, location type/name, notes, coordinates/context, analysis fields, timestamps",
    "support_tickets": "id, user_id, category, subject/description, priority, status, assignee/resolution, timestamps",
    "roles": "role identifier/name and display/permission metadata (baseline definition is not included)",
    "audit_logs": "actor, action, entity, old/new/detail payload, timestamp (baseline definition is not included)",
}


def table_description(name: str) -> str:
    descriptions = {
        "app_outbox": "Transactional side-effect outbox for attendance and related operations.",
        "attendance_action_requests": "Idempotent attendance command envelope and result record.",
        "attendance_policy_history": "Effective-dated copy of attendance policies.",
        "attendance_runtime_settings": "Operational attendance/tracking settings such as stale thresholds.",
        "biometric_import_batches": "Import batch header and idempotency state.",
        "biometric_import_batch_items": "Per-record biometric import results.",
        "biometric_reconciliation_queue": "Punches awaiting reconciliation/mapping.",
        "daily_route_summaries": "Derived per-day route distance and quality summary.",
        "employee_assignment_history": "Effective-dated office/team/reporting assignment.",
        "event_announcements": "Per-user event announcement and acknowledgement.",
        "event_attendance_policy": "Event attendance-to-HRMS behavior and duty times.",
        "event_comp_off_credits": "Legacy/auditable event comp-off credit record.",
        "event_department_assignments": "Targeted event audience by department.",
        "event_employee_assignments": "Targeted event audience by employee.",
        "event_team_assignments": "Targeted event audience by team.",
        "geocoding_cache": "Reverse-geocoding cache to limit third-party calls.",
        "holiday_calendar": "Scoped holidays and non-working days.",
        "weekly_off_rules": "Scoped recurring weekly-off rules.",
        "otp_login_challenges": "Hashed, expiring OTP challenge and attempt counter.",
        "predefined_sites": "Managed client/site/meeting destinations and geofence radius.",
        "system_admin_login_credentials": "Private static bootstrap administrator verifier; security remediation required.",
        "teams": "Normalized team master used by profiles and scoped calendar rules.",
        "retention_settings": "Database retention policy configuration.",
        "maintenance_job_runs": "Bounded maintenance/statistics job audit.",
    }
    return descriptions.get(name, f"{TABLE_DOMAINS.get(name, 'Supporting')} data object; refer to defining migration and RLS policies for the full contract.")


def add_flow(doc: Document, title: str, steps: list[tuple[str, str]]) -> None:
    doc.add_heading(title, level=3)
    add_table(doc, ["Step", "Actor / component", "Action and resulting data"],
              [(i + 1, actor, action) for i, (actor, action) in enumerate(steps)], [0.45, 1.45, 5.4], 8.5)


def main() -> None:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    props = doc.core_properties
    props.title = "BOP Realty HRMS - Complete Technical and Functional Documentation"
    props.subject = "Source-audited architecture, workflows, data, UI, security and operations"
    props.author = "OpenAI Codex - repository-derived documentation"
    props.keywords = "HRMS, Expo, React Native, Supabase, attendance, biometric, leave, field visits"

    # Cover
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo = ROOT / "assets" / "images" / "logo.png"
    if logo.exists():
        p.add_run().add_picture(str(logo), width=Inches(2.5))
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("BOP REALTY HRMS")
    r.bold = True; r.font.size = Pt(34); r.font.color.rgb = RGBColor.from_string(PURPLE_DARK)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Complete Technical, Functional, Data, Workflow, UI and Security Documentation")
    r.font.size = Pt(17); r.font.color.rgb = RGBColor.from_string(PURPLE)
    doc.add_paragraph()
    add_table(doc, ["Document field", "Value"], [
        ("Application", "Bop Realty / RE HRMS"), ("Application version", VERSION), ("Audited build ID", BUILD),
        ("Required backend schema", SCHEMA), ("Repository snapshot", TODAY),
        ("Basis", "Static audit of active source, migrations, tests, configuration and release notes"),
        ("Confidentiality", "Internal technical documentation; contains architecture and security findings but no secret values"),
    ], [1.8, 5.2], 9)
    doc.add_paragraph("This document describes what the current repository actually implements. Older project notes are treated as historical and are not assumed to be correct when they conflict with active code or later migrations.")
    doc.add_page_break()

    doc.add_heading("Document control and audit boundaries", level=1)
    add_table(doc, ["Item", "Coverage / limitation"], [
        ("Covered", "Application UI, navigation, role gates, services, client/server data flow, Edge Functions, SQL migrations/scripts, data objects, integrations, offline behavior, build/deploy, tests and security."),
        ("Authority order", "Active TypeScript and latest ordered migrations > deployment/regression reports > older implementation/design notes."),
        ("Database limitation", "The repository has incremental migrations but no complete initial production schema dump. Core pre-existing tables are documented from runtime references; exact remote columns, policy state, row counts and extensions require a live Supabase catalog export."),
        ("Security limitation", "This is a source/configuration audit, not a penetration test, mobile binary reverse-engineering exercise or live cloud review."),
        ("Secrets", "Values from .env and privileged credentials are intentionally excluded. Public variable names and exposure implications are documented."),
        ("Generated artifacts", "dist, node_modules, APK/AAB/ZIP contents and generated hosting bundles are not line-by-line source-authority; release identity and packaging are covered."),
    ], [1.55, 5.55], 8.7)
    add_note(doc, "Critical action", "Rotate and remove the repository-defined static SYSTEM_ADMIN OTP path before further production use. See Security Assessment F-01.", AMBER)

    doc.add_heading("Table of contents", level=1)
    add_toc(doc)
    doc.add_page_break()

    # 1 Executive summary
    doc.add_heading("1. Executive system overview", level=1)
    doc.add_paragraph("BOP Realty HRMS is a cross-platform, real-estate-oriented HR and field-work system. The primary client is an Expo/React Native application delivered as a web/PWA experience, iOS bundle, and an Android native capability shell that loads the hosted HRMS in a WebView with a bundled fallback. Supabase supplies PostgreSQL, authentication, storage, realtime subscriptions, Edge Functions and scheduled database jobs.")
    add_bullets(doc, [
        "Employee capabilities: OTP login, dashboard, office attendance, attendance history and regularization, background tracking, field/site/meeting visits, tasks, event attendance, leave, activity logs, notifications, support tickets, team visibility and profile/device management.",
        "HR/admin capabilities: employee lifecycle, hierarchy, attendance reconciliation/correction, biometric mapping/import/sync, calendars, leave/comp-off ledger, office/site geofences, events/audiences, task control, notifications, live tracking, fraud review, audit logs, tickets, reports and TV display.",
        "Canonical attendance combines app, biometric and verified field-visit evidence. Earliest valid first-in wins; checkout follows a protected precedence chain and deliberately treats biometric last punch as evidence rather than automatic official checkout.",
        "India business time is authoritative (Asia/Kolkata); shared server clock and canonical date functions avoid browser timezone drift.",
        "The codebase contains 53 screen files, 55 service files, 9 Edge Functions, 99 ordered migrations, 35 passing test suites and 202 passing tests in this audit.",
    ])

    doc.add_heading("1.1 High-level architecture", level=2)
    add_table(doc, ["Layer", "Implementation", "Responsibilities"], [
        ("Presentation", "React 19.2 / React Native 0.85 / Expo 56", "Responsive screens, theme, forms, maps, native permissions and reports."),
        ("Navigation", "React Navigation 7", "Auth root, conditional tabs, stack routes and deep links."),
        ("Client domain", "TypeScript services/contexts/tasks", "Workflow orchestration, geofence checks, offline queue, realtime refresh and RPC adapters."),
        ("Native capability", "Expo modules + Android WebView shell", "Location, notifications, updates, file sharing and bridge calls."),
        ("Backend API", "Supabase PostgREST/RPC + 9 Deno Edge Functions", "Protected mutations, OTP, biometric integrations, evidence upload and push dispatch."),
        ("Data", "Supabase PostgreSQL + Storage + Realtime", "HR records, attendance, field/location history, leave ledger, tasks, events, audit and evidence."),
        ("Scheduled processing", "pg_cron / Edge dispatch", "Biometric sync, notification rules, maintenance, accrual and comp-off expiry."),
        ("External services", "Google Maps/APIs, EduMarc SMS, SmartOffice, Expo Push/EAS", "Maps/geocoding, OTP delivery, biometric punches and mobile delivery."),
    ], [1.25, 2.15, 3.75], 8.4)
    add_code(doc, "User -> Expo/React Native UI -> services/contexts -> Supabase Auth + PostgREST/RPC\n"
                  "                                      -> Edge Functions -> SMS / SmartOffice / Expo Push\n"
                  "                                      -> PostgreSQL / Storage / Realtime / pg_cron\n"
                  "Android hosted UI <-> validated native bridge <-> location, notifications, updates, sharing")

    doc.add_heading("1.2 Technology stack", level=2)
    package = json.loads(read("package.json"))
    stack_rows = [("Language", "TypeScript (strict mode), SQL/PLpgSQL, Kotlin native module, PowerShell release scripts"),
                  ("Client", "Expo 56, React 19.2.3, React Native 0.85.3, React Native Web 0.21.2"),
                  ("Backend client", "@supabase/supabase-js 2.x"),
                  ("Maps/location", "expo-location, react-native-maps, Google Maps JS/Embed/Routes/Geolocation APIs"),
                  ("Offline/network", "NetInfo, SecureStore, AsyncStorage, service worker"),
                  ("Documents/import", "expo-document-picker, SheetJS xlsx 0.20.3"),
                  ("Testing/quality", "Jest 29, jest-expo, TypeScript 6, ESLint 9, Prettier 3"),
                  ("Release", "EAS Build/Update, static Metro export, Hostinger-style Apache hosting")]
    add_table(doc, ["Area", "Technology"], stack_rows, [1.5, 5.6], 8.7)

    # 2 Repository/config
    doc.add_heading("2. Repository structure and configuration", level=1)
    repo_rows = [
        ("App.tsx", "Composition root; providers, navigation, Android hosted shell, PWA/offline/update banners."),
        ("screens/", "Employee and HR/admin presentation modules."), ("navigation/", "Role-aware tabs, stacks and deep links."),
        ("components/", "Shared UI, attendance checkout, maps, native shell and modals."), ("contexts/", "Auth and location tracking lifecycle."),
        ("services/", "Domain services, Supabase access, background tracking and native bridge."), ("utils/", "Time/date, GPS, storage, formatting, error and pagination helpers."),
        ("supabase/migrations/", "Ordered production database evolution and release contracts."), ("supabase/functions/", "Deno server-side functions."),
        ("scripts/", "Historical/operational SQL, build/export/seed and conversion utilities; not all are safe for production."),
        ("public/", "PWA worker, manifest, Apache headers/SPA fallback."), ("android/", "Native Android project generated/configured for Expo."),
        ("__tests__/ + nested tests", "Regression, security-hardening, workflow and utility tests."), ("docs/", "Release/regression/backfill reports and this generated document."),
    ]
    add_table(doc, ["Path", "Purpose"], repo_rows, [1.85, 5.25], 8.5)

    doc.add_heading("2.1 Runtime configuration", level=2)
    add_table(doc, ["Variable", "Sensitivity", "Use"], [
        ("EXPO_PUBLIC_SUPABASE_URL", "Public endpoint", "Validated HTTPS Supabase project URL."),
        ("EXPO_PUBLIC_SUPABASE_ANON_KEY", "Public/publishable", "Client project key; never treat as a server secret."),
        ("EXPO_PUBLIC_GOOGLE_MAPS_API_KEY", "Public but restrict", "Native/JS/Embed maps; apply app, referrer and API restrictions."),
        ("EXPO_PUBLIC_EAS_PROJECT_ID", "Public identifier", "Expo push token and EAS project identity."),
        ("EXPO_PUBLIC_WEB_BASE_PATH", "Public", "Static host root, router and service-worker scope (for example /HRMS/)."),
        ("EXPO_PUBLIC_LOCATION_RPC_ENABLED", "Feature flag", "Enables protected server location path."),
        ("EXPO_PUBLIC_BUILD_ID", "Release metadata", "Overrides build/cache identity."),
        ("GOOGLE_SERVICES_JSON", "Build secret/file", "Firebase/FCM client configuration injected by EAS."),
        ("SUPABASE_SERVICE_ROLE_KEY", "Server secret", "Edge Function privileged database access; never enters client bundle."),
        ("SMS_OTP_API_KEY", "Server secret", "EduMarc SMS OTP delivery."),
        ("SmartOffice/Vault secrets", "Server/Vault secret", "Biometric vendor configuration retrieved server-side."),
    ], [2.35, 1.2, 3.6], 8.1)
    add_note(doc, "Configuration hygiene", ".env and .env.local are not tracked; only .env.example is tracked. Never add production values to source, ZIP packages or screenshots.")

    doc.add_heading("2.2 Release and platform configuration", level=2)
    add_bullets(doc, [
        f"Application version {VERSION}; release build ID {BUILD}; required backend schema {SCHEMA}.",
        "Expo scheme re-hrms; Android package and iOS bundle identifier are com.rehrms.app.",
        "Portrait default; iPad/tablet support enabled; automatic light/dark UI style.",
        "Android permissions: camera, coarse/fine/background location, foreground location service and notifications.",
        "iOS background modes: location and remote notification; explicit camera/location permission copy is configured.",
        "EAS channels: development, preview and production. Production normally builds an Android App Bundle; internal preview builds APKs.",
        "Runtime version follows appVersion. OTA updates are checked on launch, with immediate cached fallback.",
        "The current release config says Android versionCode 9; release notes identify an internally built auto-incremented APK as build 11. Treat EAS remote versioning as authoritative for distributed artifacts.",
    ])

    # 3 UI theme
    doc.add_heading("3. Visual design system and color theme", level=1)
    doc.add_paragraph("The active theme is purple-led and differs from the older blue design_guidelines.md concept. constants/theme.ts is authoritative.")
    light_colors = [("Primary", "#6D1B6D", "Brand/navigation/action"), ("Primary light", "#8E2A8E", "Hover/dark-theme accent"),
                    ("Primary dark", "#4A0F4A", "Headers/emphasis"), ("Root background", "#F5F0F8", "App canvas"),
                    ("Surface", "#FFFFFF", "Cards/forms"), ("Secondary background", "#F0EAF5", "Grouped areas"),
                    ("Primary text", "#1A1028", "Main copy"), ("Secondary text", "#6B6478", "Supporting copy"),
                    ("Border", "#E2D5EA", "Controls/cards"), ("Success", "#12966E", "Approved/present"),
                    ("Warning", "#D97706", "Pending/attention"), ("Error", "#C8252D", "Rejected/failure"),
                    ("Info", "#1A5FD4", "Informational state"), ("Accents", "#F59B3C / #E84C8C / #3CB4F5 / #7C3AED", "Dashboard cards")]
    table = add_table(doc, ["Token", "Hex", "Primary use"], light_colors, [1.7, 1.25, 4.1], 8.6)
    for row, (_, color, _) in zip(table.rows[1:], light_colors):
        set_cell_shading(row.cells[1], color.lstrip("#"))
        # readable swatch label
        lum = sum(int(color.lstrip("#")[i:i+2], 16) for i in (0, 2, 4))
        set_cell_text(row.cells[1], color, bold=True, color=WHITE if lum < 350 else INK, size=8.6)
    doc.add_heading("3.1 Dark palette", level=2)
    add_table(doc, ["Group", "Tokens"], [
        ("Text", "#F0EBF8 primary; #9C8FAC secondary; #6B6478 tertiary"),
        ("Brand", "#8E2A8E primary; #A838A8 light; #6D1B6D dark"),
        ("Background", "#140E1C root; #1E1530 surface; #281A3C secondary; #341F4C tertiary"),
        ("Borders", "#382248 / #281A3C"), ("Semantic", "Success #1AB88A; warning #F5B040; error #EF7070; info #6098F8")
    ], [1.4, 5.65], 8.6)
    doc.add_heading("3.2 Spacing, shape, type and elevation", level=2)
    add_table(doc, ["System", "Definition"], [
        ("Spacing", "4, 8, 12, 16, 24, 32, 48, 56 and 64 px; inputs/buttons 52; FAB 56; icon 24."),
        ("Radii", "4, 8, 12, 16, 20, 24, 28 and full/pill 9999."),
        ("Typography", "H1 28/700; H2 22/700; H3 17/600; H4 15/600; body 15/22; small 13/19; caption 11/16."),
        ("Fonts", "System fonts on native; Inter/system stack on web; platform serif/rounded/mono fallbacks."),
        ("Shadows", "Purple-tinted xs/sm/md/lg/FAB shadows; CSS box-shadow on web and elevation/shadow props on native."),
        ("TV display", "Dark #14101C canvas, #140820 top bar, #1C1028 columns; present green and absent red."),
    ], [1.4, 5.65], 8.6)
    doc.add_heading("3.3 Responsive and interaction behavior", level=2)
    add_bullets(doc, [
        "Below 768 px web width and on mobile, navigation uses a fixed bottom tab bar; wide web uses centered labeled items in a 68 px bottom rail.",
        "Screens use shared safe-area-aware scroll/flat-list/keyboard-aware containers and common themed stack headers.",
        "Deferred/lazy screen imports reduce initial web bundle work; admin/report screens load on demand.",
        "Light/dark mode follows the platform scheme and profile/settings controls; semantic tokens remain consistent.",
        "Feather icons are the primary icon set. Loading, offline, update and error boundary states are centrally handled.",
        "Accessibility labels/states exist on the custom tab controls, but a complete WCAG audit is still required; several complex tables/maps may need keyboard and screen-reader remediation.",
    ])

    # 4 roles/navigation
    doc.add_heading("4. Roles, permissions and navigation", level=1)
    add_table(doc, ["Normalized group", "Role values", "Client navigation"], [
        ("System admin", "SYSTEM_ADMIN (aliases SUPER_ADMIN, SUPERADMIN, SYSTEMADMIN, SYS_ADMIN)", "HR Panel initial tab, Reports and all privileged modules."),
        ("Admin", "ADMIN (alias ADMINISTRATOR)", "HR Panel initial tab, Reports and all privileged modules."),
        ("HR", "HR, HRD (alias HR_MANAGER)", "HR Panel initial tab and Reports."),
        ("Manager/escalation", "Business/team/sales leadership roles plus escalation_matrix membership", "Employee shell; conditional HR Panel and/or My Team. Server authorization remains decisive."),
        ("Sales manager", "SALES_MANAGER / Sales Manager", "Employee shell with My Team when reports are assigned. Direct reports are limited to Sales Executives, Telecallers and Agents."),
        ("Sales", "SALES_EXECUTIVE, BDE, SENIOR_BDE, BDM variants, SALES, TELECALLER, CO-ORDINATOR and specified associate sales roles", "Employee tabs; HR Panel and My Team hidden."),
        ("Other employee", "Remaining active roles", "Home, Tasks, Profile; My Team shown only with reports or escalation membership."),
    ], [1.3, 2.75, 3.1], 8.1)
    add_note(doc, "Authorization rule", "Navigation hiding is usability only. Protected RPCs, RLS, Edge Function JWT checks and public.is_profile_admin() are the actual security boundary.")
    doc.add_heading("4.1 Route map and deep links", level=2)
    add_table(doc, ["Area", "Routes"], [
        ("Root", "Unauthenticated -> Login; authenticated -> Main."),
        ("Employee Home", "/, /leave, /notifications, /birthdays, /activity/new, /activity, /field-visits/:id?, /support/new, /my-team, /attendance, /tracking-setup"),
        ("Tasks", "/tasks/:eventId?"), ("Account", "/account/profile, /account/leave, /account/notifications, /account/team"),
        ("HR Panel", "/admin plus employees, attendance, leaves, tasks-projects, notifications, offices, sites, audit, activity-analysis, field-visits, live-tracking, teams, tickets, events, check-in-logs, tv, leave-balance, biometric-sync, employee-import, imported-data, employee-match, suspicious-attendance, attendance-calendar, regularization, attendance-report"),
        ("Reports", "/reports plus attendance, advanced-attendance, field-visits, leaves, regularization, biometric, activity, tracking, gps-suspicious, downloads"),
        ("Native", "re-hrms:// equivalent deep-link prefix."),
    ], [1.4, 5.65], 8.1)

    # 5 auth
    doc.add_heading("5. Authentication, session and device security", level=1)
    add_flow(doc, "5.1 Standard employee OTP login", [
        ("User", "Enters a normalized Indian 10-digit phone number."),
        ("Client/AuthContext", "Invokes otp-auth action=request; unknown numbers receive a generic success response."),
        ("otp-auth", "Allows at most 3 challenges for the phone in 10 minutes, verifies an active profile, generates a random six-digit OTP, salts/hashes it, stores a five-minute challenge and sends via EduMarc SMS."),
        ("User", "Enters the OTP; the server permits at most five failed attempts for the current challenge."),
        ("otp-auth", "Consumes the challenge, validates the active profile, creates/locates a matching Supabase Auth user and returns a magic-link token hash."),
        ("AuthContext", "Calls Supabase verifyOtp with the token hash, stores the session, fetches authoritative profile data and registers the device session/push token."),
        ("Navigation", "Mounts employee or HR shell from normalized authoritative role."),
    ])
    doc.add_heading("5.2 Session persistence and logout", level=2)
    add_bullets(doc, [
        "Native sessions/profile cache use Expo SecureStore with device-only accessibility; older AsyncStorage values are migrated and deleted.",
        "Web authentication persists in localStorage to survive browser closure; non-auth sensitive queues use tab sessionStorage and are prohibited for web location evidence.",
        "Supabase auto-refreshes tokens and does not parse sessions from URL fragments.",
        "On logout the app deactivates the push token, closes device session, stops native tracking, clears offline location/evidence queues, clears cached user metadata and signs out of Supabase.",
        "App foreground/resume logic refreshes sessions and authoritative profile state, preventing a stale cached employee shell from limiting or granting the wrong navigation.",
    ])
    add_note(doc, "Bootstrap exception", "The SYSTEM_ADMIN path uses a repository-defined static credential verifier and must be removed/rotated; it is not equivalent to the random five-minute employee OTP flow.", AMBER)

    # 6 employee workflows
    doc.add_heading("6. Employee workflows and business logic", level=1)
    doc.add_heading("6.1 Dashboard", level=2)
    add_bullets(doc, [
        "Loads a cached rpc_employee_dashboard_summary, today’s attendance state, unread notifications, tasks, leave balance, visits, birthdays/events and tracking state.",
        "Realtime attendance changes invalidate and reload relevant dashboard state. Cache invalidation follows attendance mutations.",
        "Quick actions route to attendance, leave, activity, field visits, tasks, team and tickets. HR users are redirected/mounted into the HR Panel by role.",
        "Offline banner explicitly prevents users from assuming attendance actions are committed without server confirmation.",
    ])
    add_flow(doc, "6.2 Office attendance check-in", [
        ("User/device", "Requests precise foreground location; UI rejects missing/poor location and gathers device integrity signal."),
        ("AttendanceService", "Loads active/allowed offices, selects nearest office and validates configured radius (10-1,000 m)."),
        ("Client geofence", "Rejects accuracy over 100 m; requires distance + accuracy <= configured radius."),
        ("Protected RPC", "rpc_attendance_checkin_v3 validates identity, active office/policy, server-side geofence and idempotent attendance action."),
        ("Database", "Creates/updates one employee/day attendance record, app evidence, audit/outbox side effects and committed timestamp."),
        ("Tracking lifecycle", "Starts/ensures an attendance tracking session on native, even without a field-visit destination."),
        ("UI", "Only shows success after a committed server timestamp; summaries/realtime channels refresh."),
    ])
    add_flow(doc, "6.3 Checkout, early checkout and undo", [
        ("User", "Chooses checkout; location, reason and optional notes are captured."),
        ("Protected reconciled RPC", "Validates the record, server time, policy minimums/earliest checkout and checkout geofence/context."),
        ("Early checkout", "Where required, a reason/notes flow is used; error codes are mapped to actionable UI."),
        ("Database", "Commits checkout and its audit together; protection triggers reject unaudited direct current-day checkout edits."),
        ("Tracking", "Stops only after the authoritative checkout exists."),
        ("Undo", "rpc_attendance_undo_checkout is available within the policy window (documented default five minutes)."),
    ])
    doc.add_heading("6.4 Canonical attendance calculation", level=2)
    add_table(doc, ["Decision", "Rule"], [
        ("Business date/time", "Asia/Kolkata server date and hrms_timestamp compatibility handling for legacy offset-less values."),
        ("First-in candidates", "App first-in, earliest biometric first punch and earliest verified field-visit check-in."),
        ("Official first-in", "Earliest non-null candidate; source is APP, BIOMETRIC or FIELD_VISIT. Multiple evidence is marked MULTIPLE."),
        ("Checkout precedence", "Approved regularization > admin corrected > app checkout > field-visit checkout only when policy explicitly allows it."),
        ("Biometric last punch", "Evidence only; it does not become official checkout by itself."),
        ("Late", "Official first-in at/after policy late time; default policy backfill uses 10:32 IST where office default is absent."),
        ("No configuration", "Visible CONFIG_REQUIRED status rather than disappearing or failing the page."),
        ("Calendar", "Holiday/weekly-off resolution is scoped employee > team > office > company and combined with approved leave."),
        ("Warnings", "Configuration, GPS/tracking, evidence mismatch, stale ping and missing checkout are surfaced separately."),
    ], [1.7, 5.35], 8.3)
    doc.add_heading("6.5 Attendance history and regularization", level=2)
    add_bullets(doc, [
        "Employee month view uses rpc_get_my_monthly_attendance_v3; raw evidence and official source are displayed without silently replacing canonical decisions.",
        "A regularization request supplies requested in/out values and reason. The server owns the user identity and request status.",
        "HR lists pending/approved/rejected requests and reviews through protected RPC. Approval becomes the highest checkout precedence and is audited transactionally.",
    ])

    add_flow(doc, "6.6 Field visit and current-location fallback", [
        ("Employee", "Creates an OFFICE, SITE or MEETING visit with planned destination; source journey origin is preserved."),
        ("Service/RPC", "rpc_mutate_own_field_visit creates or changes only the caller’s visit; duplicate/live-origin repair logic filters replaced rows."),
        ("Journey", "IN_TRANSIT state activates destination context and adaptive location updates; navigation can open Google Maps."),
        ("Arrival", "Selected-site check-in uses server geofence. Current-location fallback can capture precise position plus required/optional photo evidence."),
        ("Evidence upload", "JWT-authenticated Edge Function accepts only JPEG/PNG <=2 MB, stores a user-prefixed private object, and supports safe orphan deletion."),
        ("Evidence submit", "submit_field_visit_current_location_v3 binds the authenticated user, idempotency key, integrity material, capture time and evidence path."),
        ("Reviewer", "Authorized reviewer accepts/rejects fallback evidence; audit data is immutable to ordinary clients."),
        ("Checkout", "Protected own-visit mutation records completion; policy may optionally allow verified field checkout to supply official attendance checkout."),
    ])
    doc.add_heading("6.7 Background location lifecycle", level=2)
    add_bullets(doc, [
        "Native tracking starts from committed attendance and runs as an Android foreground/background location task; web clearly cannot promise continuous background tracking.",
        "Adaptive sampling intervals are 12/30/60/120 seconds depending on movement/proximity/health context; exact choice is coordinated in location services.",
        "Pings use client event IDs for idempotency, server-side identity and session validation, bounded accuracy/speed/time checks and journey distance state.",
        "Up to 200 offline pings are stored as separate SecureStore entries on native; oldest entries are dropped when the cap is reached; web queueing is disabled.",
        "Reconnect and app state transitions flush in capture-time order. Health events record permissions, GPS, stale ping, battery/session/bridge/server rejection diagnostics.",
        "Realtime attendance/visit channels update tracking lifecycle; logout clears encrypted queues and stops tracking.",
    ])

    doc.add_heading("6.8 Leave and comp-off", level=2)
    add_bullets(doc, [
        "Employees view sick/casual/earned/comp-off balances, scoped work calendar and request history; submit/cancel operations are protected RPCs.",
        "Working days exclude resolved holidays/weekly offs. Approval/debit is server-side and locks the balance row.",
        "The immutable leave_transactions ledger uses idempotency keys, records before/after balance and rejects negative balances.",
        "One casual leave unit accrues per active non-service employee per calendar month, including activation month, via idempotent scheduled processing.",
        "Comp-off earnings require source/reason, approval and optional expiry. Consumption uses oldest expiring approved units first; expiry posts a compensating ledger entry.",
        "HR manual adjustment requires a reason and client UUID idempotency key. Direct browser writes to ledger, earnings and balance are revoked by the latest migration.",
    ])

    doc.add_heading("6.9 Tasks, events, activities, notifications and support", level=2)
    add_table(doc, ["Workflow", "Behavior"], [
        ("Tasks", "Employee loads assigned tasks, filters/views details and updates status. Admin creates/assigns/controls tasks. Database workflow triggers notifications; event IDs can deep-link into Tasks."),
        ("Events", "Employee receives targeted announcement, acknowledges it, and can check in/out within event rules. Audience may be all, team, department or employee. Event attendance can synchronize to HRMS attendance and approved marketing/event rules can award comp-off."),
        ("Activities", "Employee logs activity type, location and notes; client analysis extracts entities/classification and flags weak entries. Admin reporting joins attendance-day context and exports analytics."),
        ("Notifications", "In-app rows plus unified push outbox/delivery records. Deep links open events/tasks/field visits/permissions. Native categories provide acknowledgement/navigation actions."),
        ("Support tickets", "Employee creates and tracks tickets; HR/admin triages and closes. WhatsApp/SMS client delivery is intentionally disabled until a protected server provider exists."),
        ("Profile and device", "Employee edits allowed profile/avatar fields, sees office/reporting data, lists/revokes device sessions, changes theme and logs out."),
        ("My Team", "Eligible managers see direct reports, attendance/task/visit status and latest authorized ping; database RLS/manager relationship must enforce the same scope."),
    ], [1.35, 5.7], 8.2)

    # 7 admin
    doc.add_heading("7. HR and administrator workflows", level=1)
    admin_rows = [
        ("HR dashboard", "rpc_admin_dashboard_summary provides counts; cards route to every control/report module."),
        ("Employee management", "Protected profile save/import/activation RPCs; office/role validation; bulk employee profiles; service/TV/reporting flags."),
        ("Hierarchy", "Protected v2 hierarchy update and transactional batch import; reporting cycles/invalid managers are rejected server-side. Sales Manager is a first-class position whose direct reports are restricted to Sales Executives, Telecallers and Agents."),
        ("Attendance control", "Canonical V3 employee-day list, app/biometric/field evidence, warnings, live state, corrections, biometric mapping and exports."),
        ("Office/site config", "Protected configuration mutation; active office geofence/radius validation and configuration issue report; predefined site review."),
        ("Calendar", "Scoped holiday and weekly-off rules with protection against past holiday rewrites."),
        ("Leave", "Review requests, inspect calendar, grant comp-off, view ledger summary and make idempotent balance adjustments."),
        ("Biometric", "SmartOffice connection via Vault-backed proxy, secure sync, XLS import, employee-code mapping/unmapping, unmatched punch reprocessing and reconciliation."),
        ("Field/tracking", "Field visit report/evidence review, live tracking map, explicit ping request, alert resolution, suspicious attendance review and TV occupancy."),
        ("Events", "Create targeted events with venue/radius/duty times/audience, inspect attendees, delete via protected RPC and optionally attach event photo."),
        ("Tasks/projects", "Create, assign, prioritize and monitor employee tasks; 'Projects' is task/project-style control, not the inactive real-estate project directory screen."),
        ("Notifications", "Create per-user in-app notifications; push rules/outbox deliver on native devices."),
        ("Audit and operations", "Audit log views, check-in attempts, client operational errors, migration ledger, device sessions and maintenance job records."),
        ("Reports", "Dedicated attendance, advanced, field, leave, regularization, biometric, activity, tracking, GPS/suspicious and download routes."),
    ]
    add_table(doc, ["Module", "Workflow and controls"], admin_rows, [1.55, 5.5], 8.2)
    add_note(doc, "Privilege design", "Latest protected mutations revoke anon access and perform role checks in SQL/Edge Functions. Historical scripts that disable RLS are not safe operational procedures.", AMBER)

    # 8 data flow
    doc.add_heading("8. End-to-end data flow and state ownership", level=1)
    add_table(doc, ["Data class", "Authoritative owner", "Client behavior", "Propagation"], [
        ("Identity/role", "profiles + Supabase Auth", "Caches for startup, then refreshes authoritative profile.", "AuthContext -> navigation remount."),
        ("Attendance", "Protected RPCs + attendance_logs + canonical view", "Never treats local button completion as committed state.", "Outbox/triggers, realtime, dashboard/report reload."),
        ("Biometric", "SmartOffice logs/punch tables and reconciliation functions", "Admin initiates sync/import/mapping.", "Canonical first-in and admin reconciliation views."),
        ("Location", "tracking sessions + location_tracking", "Native queues signed/idempotent pings; web is foreground-only.", "Realtime maps, route summaries, alerts and health events."),
        ("Field evidence", "Private Storage + field_visits/audit", "Compress/upload, then protected metadata submit; retry is idempotent.", "Reviewer decision and canonical field evidence."),
        ("Leave", "leave_balance + immutable transactions + requests", "Reads balance/calendar; writes through RPC only.", "Approval transaction and audit/notification."),
        ("Notifications", "notifications + notification_push_outbox", "Reads/acknowledges/deep-links.", "Database rules -> outbox -> Edge dispatcher -> Expo receipts."),
        ("Configuration", "offices/sites/policies/calendar/settings", "Admin edits through protected RPC where implemented.", "Effective-dated/canonical queries consume configuration."),
    ], [1.2, 1.9, 2.05, 2.0], 7.8)
    doc.add_heading("8.1 Error, retry and idempotency conventions", level=2)
    add_bullets(doc, [
        "Mutation services verify success payloads and committed identifiers/timestamps instead of accepting a transport-level 200 as success.",
        "Attendance action requests, import batches, field evidence, location pings, profile imports, hierarchy imports and leave adjustments have explicit idempotency concepts.",
        "Offline attendance buttons remain unavailable without server confirmation; only location/evidence queues have bounded retry behavior.",
        "Operational errors may be sent to client_operational_errors through a protected RPC; production Babel removes console statements from bundled client output.",
        "Business-clock RPC and cached date prevent local clock manipulation from deciding HR events.",
    ])

    # 9 service/edge
    doc.add_heading("9. Client service and API catalogue", level=1)
    services = collect_services()
    add_table(doc, ["Service file", "Purpose", "Direct data/API dependencies"], [
        (s["path"], s["description"], "; ".join((["tables: " + ", ".join(s["tables"])] if s["tables"] else []) +
                                                      (["RPC: " + ", ".join(s["rpcs"])] if s["rpcs"] else []) +
                                                      (["Edge: " + ", ".join(s["edge"])] if s["edge"] else [])) or "Local/platform logic")
        for s in services
    ], [2.05, 3.0, 2.15], 7.0)

    doc.add_heading("9.1 Supabase Edge Functions", level=2)
    edge_rows = [
        ("otp-auth", "Public POST request/verify; rate-limited challenge; service-role profile/Auth administration; SMS provider.", "Random employee OTP; special static admin path is unsafe."),
        ("smartoffice-proxy", "Authenticated HR/admin proxy; retrieves vendor secrets from Database Vault.", "Prevents SmartOffice credentials entering the client."),
        ("biometric-attendance-sync", "Scheduled project-key or authenticated admin invocation; claims sync and writes vendor punch data.", "Idempotent sync/release contract."),
        ("biometric-manual-import", "JWT-authenticated admin-only manual punch import.", "Server validates role and imports protected batch."),
        ("biometric-punch-webhook", "Shared-secret vendor webhook.", "Ingests individual SmartOffice punches."),
        ("field-visit-evidence-upload", "JWT-authenticated multipart upload/orphan delete.", "Private JPEG/PNG <=2 MB; owner-prefixed path."),
        ("notification-dispatch", "Scheduled push outbox/receipt processor.", "Currently authenticates with publishable anon key pair; remediate."),
        ("send-journey-notifications", "Shared-secret journey notification handler.", "Manager/employee journey state notifications."),
        ("send-event-notifications", "Event notification sender using service-role state.", "Legacy/specialized alongside unified outbox; deployment status must be verified."),
    ]
    add_table(doc, ["Function", "Purpose/authentication", "Notes"], edge_rows, [1.75, 3.35, 2.1], 7.7)

    # 10 DB
    doc.add_heading("10. Database architecture and complete object catalogue", level=1)
    objects = collect_sql_objects()
    ts_paths = list((ROOT / "screens").rglob("*.tsx")) + list((ROOT / "services").rglob("*.ts")) + list((ROOT / "contexts").rglob("*.tsx")) + list((ROOT / "components").rglob("*.tsx"))
    runtime_tables, runtime_rpcs, runtime_edges = ts_refs(ts_paths)
    all_tables = sorted(set(objects["table"]) | set(runtime_tables))
    doc.add_paragraph(f"The repository defines or references {len(all_tables)} unique table/storage names, {len(objects['view'])} SQL views, {len(objects['function'])} SQL functions and {len(objects['policy'])} named policies across migrations and operational scripts. Counts include legacy generations; the latest migration definition/grant order controls production behavior.")
    add_note(doc, "Reconstruction caveat", "Core tables such as profiles, attendance_logs, offices, tasks, notifications, leave_balance, roles and audit_logs predate the included migrations. Their columns below are runtime-inferred unless a definition source is listed. Obtain pg_dump --schema-only and pg_catalog policy/grant exports for an exact live data dictionary.")
    table_rows = []
    for name in all_tables:
        sources = sorted(objects["table"].get(name, set()))
        consumers = sorted(runtime_tables.get(name, set()))
        columns = CORE_COLUMNS.get(name, "See defining migration; columns are intentionally not guessed in this catalogue.")
        evidence = ("DDL: " + ", ".join(sources[-2:])) if sources else "Runtime referenced; baseline DDL absent"
        if consumers:
            evidence += "; used by " + ", ".join(consumers[:3]) + ("…" if len(consumers) > 3 else "")
        table_rows.append((name, TABLE_DOMAINS.get(name, "Supporting / audit / integration"), table_description(name), columns, evidence))
    add_table(doc, ["Table / bucket", "Domain", "Purpose", "Key columns / contract", "Definition / usage evidence"], table_rows,
              [1.25, 1.2, 2.0, 2.4, 2.1], 6.3)

    doc.add_heading("10.1 Views and derived reporting contracts", level=2)
    view_desc = {
        "v_attendance_daily_canonical": "Employee/date spine with raw evidence, official first-in/last-out, source, calendar, leave, warning and tracking health.",
        "attendance_reconciliation": "Biometric/app reconciliation compatibility view; latest definitions supersede earlier generations.",
        "attendance_reconciliation_v2": "Operational reconciliation/report view for current-day and admin inspection.",
        "employee_live_status": "Legacy/admin live employee attendance and tracking status.",
        "event_admin_summary": "Event audience/attendance counts and administrative event facts.",
        "admin_attendance_export": "Legacy export view.", "admin_daily_employee_log": "Legacy joined employee-day activity/attendance view.",
    }
    add_table(doc, ["View", "Purpose", "Definition sources"], [(name, view_desc.get(name, "Derived SQL view."), ", ".join(sorted(paths))) for name, paths in sorted(objects["view"].items())], [1.8, 3.65, 1.85], 7.5)

    doc.add_heading("10.2 Function/RPC catalogue", level=2)
    function_rows = []
    for name, paths in sorted(objects["function"].items()):
        category = ("Client RPC" if name.startswith("rpc_") else "Trigger/scheduled/helper")
        consumers = sorted(runtime_rpcs.get(name, set()))
        purpose = name.replace("rpc_", "").replace("_", " ").capitalize()
        if consumers:
            purpose += "; called by " + ", ".join(consumers[:3])
        function_rows.append((name, category, purpose, ", ".join(sorted(paths)[-3:])))
    add_table(doc, ["Function", "Class", "Purpose / runtime consumer", "Latest source evidence"], function_rows, [1.85, 1.1, 2.75, 2.3], 6.4)

    doc.add_heading("10.3 RLS, grants, storage, realtime and schedules", level=2)
    add_bullets(doc, [
        "Latest migrations enable RLS for tracking, field evidence metadata, imports, biometric, event audience, device sessions, leave ledger, outbox and audit/diagnostic tables; public.is_profile_admin() is a shared privileged predicate.",
        "Protected RPCs generally revoke public/anon execute, grant authenticated, derive the acting identity from auth.uid() and perform explicit admin/ownership checks.",
        "A security_audit.sql checklist verifies RLS/policies, anon grants, executable functions and private storage. It must be run against each deployed environment; repository state cannot prove remote state.",
        "Private storage buckets defined in migrations: field-visit-evidence (2 MB JPEG/PNG) and profile-pics. The Event screen also assumes event-photos and obtains a public URL, but no bucket migration was found.",
        "Realtime subscriptions watch attendance_logs, location_tracking, field_visits, location_alerts and smartoffice_device_logs for employee, team, admin and TV refresh.",
        "Current schedules include minute-level notification rule evaluation/dispatch, biometric synchronization/reconciliation generations, hourly/daily maintenance/statistics, monthly casual leave accrual and daily comp-off expiry. Later migrations unschedule/replace earlier job generations.",
        "Retention tables and bounded maintenance functions exist for punch/location/route/operational data; actual configured retention values must be queried from production.",
    ])
    doc.add_heading("10.4 Data integrity and audit controls", level=2)
    add_bullets(doc, [
        "One employee/date attendance record is protected by duplicate cleanup/unique guardrails and transactional action request patterns.",
        "Current-day checkout changes require a matching HRMS audit event or explicitly authorized regularization path.",
        "Biometric import uses batches, per-item state, source hashes and rollback/reprocessing helpers.",
        "Field evidence v3 confirms authenticated identity, storage ownership and idempotency while raw audit rows reject mutation.",
        "Leave ledger changes lock balances, record before/after state and use unique idempotency keys; compatibility triggers audit legacy balance mutations.",
        "Migration/backfill operations preserve timestamp backups and record REVIEW_REQUIRED items instead of silently forcing unsafe conversions.",
    ])

    # 11 integrations/offline/deployment
    doc.add_heading("11. External integrations", level=1)
    add_table(doc, ["Integration", "Data exchanged", "Security/operational notes"], [
        ("Supabase", "Sessions, profiles, HR data, RPCs, realtime, storage and Edge Functions.", "HTTPS enforced in client config; publishable key is public; service role server-only."),
        ("EduMarc SMS", "Phone and OTP message.", "Server-only API key. Add provider monitoring, DLT/template controls, abuse/cost alerts and stronger global rate limits."),
        ("SmartOffice", "Employee codes, punch logs and attendance sync.", "Credentials should remain in Database Vault/Edge Function; mapping and batch audit are protected."),
        ("Google Maps", "Coordinates, map tiles/embed, navigation, geocoding/routes.", "Key is public in client by design; restrict per Android/iOS/web app and to required APIs. Location is sensitive personal data."),
        ("OpenStreetMap Nominatim", "Coordinates for reverse geocoding fallback.", "Respect rate/user-agent policy and avoid sending unnecessary identity."),
        ("Expo Push", "Expo push token, message content, receipt IDs.", "Tokens are device-scoped and deactivated on logout; avoid sensitive HR detail in lock-screen text."),
        ("Expo EAS/Updates", "Build metadata and OTA JS bundles.", "Runtime/schema contract prevents incompatible clients; channels separate production/preview."),
        ("Apache/Hostinger static web", "Single-page app, hashed bundles, manifest and service worker.", "CSP/HSTS/security headers present; correct cache headers are release-critical."),
    ], [1.4, 2.4, 3.25], 8.0)

    doc.add_heading("12. Offline, reliability and performance", level=1)
    add_bullets(doc, [
        "PWA uses network-first same-origin fetch with release-specific cache, retains current and one prior release cache, caps entries at 120 and falls back to cached shell/assets offline.",
        "Native Android WebView disables its PWA worker/caches and WebView caching so it follows hosted deployment; on load/network/HTTP failure it opens the bundled HRMS.",
        "Offline attendance mutations are disallowed because they require server confirmation. Native location and current-location evidence use bounded encrypted retry queues.",
        "Admin screens are lazy loaded; dashboard summaries have short caches; Supabase realtime is limited to 10 events/second on the client.",
        "Large admin screens (attendance control ~3,265 lines, advanced report ~1,470, employee import ~1,459, live tracking ~1,448) are maintainability/performance hotspots and should be decomposed.",
        "Pagination utilities exist and several RPCs enforce limits, but some direct .select() admin screens may depend on Supabase default row caps and need explicit pagination verification.",
        "Service-worker HTML-as-asset detection prevents a static-host SPA fallback from poisoning JavaScript cache entries.",
    ])

    doc.add_heading("13. Build, deployment and operations", level=1)
    add_table(doc, ["Operation", "Command / process", "Result"], [
        ("Local", "npm start / npm run web / npm run android", "Expo development server."),
        ("Quality", "npm test; npm run typecheck; npm run lint", "Jest, strict TypeScript and Expo ESLint."),
        ("Web export", "npm run export:web", "Static dist, versioned worker/precache and hosting package preparation."),
        ("Web deploy", "npm run deploy:web or upload bop-hrms-deploy.zip", "Deploy under configured /HRMS base path; verify cache headers and deep routes."),
        ("OTA", "npm run update:production / update:installed", "Publishes runtime-compatible Expo update to channel."),
        ("Android", "EAS production build; internal APK as needed", "Store AAB or internal APK; versioning uses EAS remote auto-increment."),
        ("Backend", "supabase db push and function deploy", "Apply migrations in order, deploy version-compatible functions, then run security audit/release probes."),
    ], [1.25, 2.6, 3.2], 8.2)
    doc.add_heading("13.1 Deployment order and release gate", level=2)
    add_numbered(doc, [
        "Back up production and review the migration/backfill report; never run destructive historical reset/disable-RLS scripts.",
        "Apply ordered migrations and confirm local/remote migration versions match.",
        "Run supabase/security_audit.sql and release contract/version probes; resolve any RLS/grant mismatch.",
        "Deploy Edge Functions with server/Vault secrets and confirm authentication behavior.",
        "Run typecheck, lint, Jest and production web export; validate no secret entered the bundles.",
        "Upload the complete static package atomically; verify index/sw no-cache and immutable hashed asset behavior.",
        "Publish compatible OTA updates and build a new binary if runtime/native configuration changed.",
        "Perform physical-device smoke tests for OTP, office attendance, background tracking, field evidence, event push and checkout.",
    ])
    doc.add_heading("13.2 Backup, monitoring and rollback", level=2)
    add_bullets(doc, [
        "Use managed PostgreSQL point-in-time recovery and scheduled logical backups; separately verify private Storage recovery and vendor configuration/Vault backup procedures.",
        "Monitor failed OTP/SMS, biometric sync state, reconciliation queue, push outbox/deliveries, client operational errors, tracking health, fraud events, cron job runs and migration report REVIEW_REQUIRED rows.",
        "OTA rollback republishes the last compatible update. Web rollback restores the previous full /HRMS archive. Database rollback should use corrective forward migrations because migrations are additive/data-preserving.",
        "Define alert thresholds and owners; the repository does not contain a configured third-party APM/error-monitoring connector.",
    ])

    # 14 testing
    doc.add_heading("14. Testing and quality status", level=1)
    add_table(doc, ["Check executed on 10 Aug 2026", "Result"], [
        ("npm test -- --runInBand", "PASS: 35 suites, 202 tests, 0 failures, 0 snapshots."),
        ("npm run typecheck", "PASS: strict tsc --noEmit."),
        ("npm run lint", "PASS with 0 errors and 1 warning: missing fallbackDestination dependency in ActiveJourneyCard.web.tsx useEffect."),
        ("npm audit --json", "27 advisories: 18 high, 9 moderate, 0 critical; 1,151 total dependency nodes in audit metadata."),
    ], [2.7, 4.35], 8.5)
    add_bullets(doc, [
        "Coverage is strongest for attendance workflows/canonical V3, security hardening contracts, auth persistence, service worker, offline location, GPS/time formatting, notifications, native bridge protocol and field event repairs.",
        "No end-to-end browser/mobile automation, coverage threshold, database integration test runner, load test or accessibility test is configured in package.json.",
        "Edge Functions are excluded from the main TypeScript config; Deno type/lint checks must be part of backend CI separately.",
        "Production release notes record successful web export and physical artifacts, but current live deployment state must be reverified independently.",
    ])

    # 15 security
    doc.add_heading("15. Security and vulnerability assessment", level=1)
    doc.add_paragraph("Risk ratings combine exploitability, privilege/data impact and whether the issue is in an active path. They are source-audit findings, not proof of live exploitation.")
    findings = [
        ("F-01", "CRITICAL", "Static bootstrap SYSTEM_ADMIN OTP is offline-recoverable", "otp-auth source fixes a privileged phone identity; a migration stores the salt and SHA-256 hash of a six-digit static OTP. Anyone with the repository can brute-force the one-million-value space offline.", "Immediately rotate/disable the credential; remove static bypass and repository verifier; use random short-lived OTP or phishing-resistant admin MFA; invalidate sessions and audit privileged actions."),
        ("F-02", "HIGH", "notification-dispatch treats the public anon key as a scheduler secret", "The function accepts requests when Authorization and apikey equal SUPABASE_ANON_KEY, which is deliberately present in every client bundle.", "Require a distinct high-entropy scheduler secret, service-to-service JWT/cron database call, or service-role-authenticated internal invocation; rate limit and audit."),
        ("F-03", "HIGH", "Embedded team map WebView permits wildcard origins and mixed content", "MyTeamScreen builds HTML with employee name interpolation that escapes only quotes, enables mixedContentMode=always and originWhitelist=['*']. A crafted name containing script-breaking characters can alter inline JavaScript.", "Use react-native-maps or a local static document with strict escaping via JSON.stringify; disallow mixed content, restrict navigation/origins and disable unnecessary DOM storage."),
        ("F-04", "HIGH (operational)", "Repository includes scripts that disable RLS on HR data", "Multiple scripts disable RLS for profiles, attendance, locations, leaves, tasks, events and audit tables. They appear historical, but accidental execution could expose data.", "Move to a quarantined archive, add irreversible warnings/guards, prohibit in runbooks/CI and verify deployed RLS with security_audit.sql on every release."),
        ("F-05", "HIGH/MODERATE", "Known dependency advisories", "npm audit reports 18 high and 9 moderate advisories across Expo/Metro/React Native and transitive packages, including DoS issues in brace-expansion, image-size, js-yaml and nanoid. No critical advisory was reported.", "Use Expo-compatible upgrade guidance; update lockfile/transitives, rerun expo-doctor/test/export, distinguish build-tool-only paths from shipped runtime, and document accepted residual risks."),
        ("F-06", "MODERATE", "Web auth tokens persist in localStorage", "PersistentAuthStorage intentionally stores the Supabase session in localStorage, increasing theft impact if any same-origin XSS occurs.", "Prefer secure HttpOnly same-site server sessions where feasible, or reduce persistence/TTL, maintain strict CSP, eliminate injection sinks and add session/device revocation monitoring."),
        ("F-07", "MODERATE", "Wildcard CORS on several Edge Functions", "otp-auth, SmartOffice, evidence upload and biometric functions return Access-Control-Allow-Origin: *. Authentication still protects most actions, but any site can originate calls and amplify OTP/scheduler abuse.", "Allow only production/approved origins where browser access is required; keep authentication and CSRF-independent bearer semantics; add rate limits."),
        ("F-08", "MODERATE", "Event photo storage is not reproducibly provisioned", "EventManagementScreen requests a public URL from event-photos, but no repository migration creates/configures that bucket or policies.", "Add a migration with explicit public/private choice, MIME/size limits, ownership/admin policies and deletion/retention behavior; prefer signed URLs for HR data."),
        ("F-09", "MODERATE", "OTP abuse controls are phone-scoped only", "Request throttling is 3 per phone per 10 minutes; IP is recorded but not enforced. There is no CAPTCHA/device/global/provider cost limit in source.", "Add IP/device/global sliding windows, provider spend alarms and adaptive challenge/CAPTCHA while retaining non-enumerating responses."),
        ("F-10", "MODERATE", "Location privacy/retention requires operational proof", "The system collects continuous precise employee locations, routes, accuracy, device/push identifiers and photo evidence. Retention machinery exists, but configured values and consent/legal controls are not provable from source.", "Document lawful basis, notice/consent, role-limited access, retention/deletion schedules, DSAR process and periodic access review; minimize lock-screen/location detail."),
        ("F-11", "LOW/MODERATE", "Android shell allows file access", "The main WebView has allowFileAccess enabled and a powerful native bridge. Origin checks restrict bridge messages to the allowed HRMS URL, which materially lowers risk.", "Disable file access if unnecessary, maintain exact-origin checks, validate navigation redirects and keep file sharing MIME/size/path constrained."),
        ("F-12", "LOW", "Development Mappls proxy has open CORS", "Metro development middleware proxies Mappls endpoints and returns wildcard origin. It is not used by static production export.", "Bind development servers to trusted interfaces, never expose Metro publicly and remove unused proxy code if Mappls is no longer active."),
        ("F-13", "LOW", "One React hook dependency warning", "ActiveJourneyCard.web.tsx omits fallbackDestination from a useEffect dependency list; stale destination state is possible.", "Include the dependency or memoize/document the invariant; add regression coverage."),
        ("F-14", "INFO", "No hardcoded service-role key detected in tracked application source", "The tracked .env files are excluded and server functions read privileged keys from environment/Vault. Older documentation contains sensitive-key terminology but not necessarily live values.", "Continue secret scanning; rotate on any accidental disclosure; exclude generated packages and logs from public distribution."),
    ]
    add_table(doc, ["ID", "Severity", "Finding", "Evidence / impact", "Required remediation"], findings, [0.45, 0.85, 1.55, 2.5, 2.45], 6.7)
    doc.add_heading("15.1 Existing security strengths", level=2)
    add_bullets(doc, [
        "HTTPS-only Supabase URL validation; service-role and SMS/vendor credentials remain server-side.",
        "Hashed, salted, expiring standard OTP challenges with non-enumerating unknown-user response, request throttling, attempt lock and consumption.",
        "Private field evidence and profile picture buckets with ownership/reviewer policies; evidence content signature and size checks.",
        "Native SecureStore, encrypted offline ping entries, logout queue clearing and device-session revocation.",
        "Protected transactional RPCs, RLS/grant hardening, server-owned identity and audit/outbox patterns.",
        "Web CSP, HSTS, nosniff, frame denial, referrer and permissions policies; production console removal.",
        "WebView native bridge validates message size/schema and allowed hosted origin before executing capabilities.",
    ])
    doc.add_heading("15.2 Prioritized remediation plan", level=2)
    add_table(doc, ["Priority", "Target", "Exit condition"], [
        ("P0 - immediate", "F-01 static admin OTP; F-02 public scheduler authentication", "Credentials rotated, static code removed, sessions audited, distinct scheduler authentication deployed and tested."),
        ("P1 - next release", "F-03 WebView injection/mixed content; verify all production RLS and grants", "Safe map implementation, strict WebView policy, security_audit.sql clean in production."),
        ("P1 - next release", "Dependency remediation", "Expo-supported dependency set with materially reduced/accepted audit findings; regression/export pass."),
        ("P2", "Web session risk, CORS, OTP abuse, event storage", "Threat-modeled decisions implemented; bucket/config migrations reproducible."),
        ("P2/P3", "Location privacy, observability, accessibility and large-file decomposition", "Approved retention/access policy, monitored controls and owned engineering plan."),
    ], [1.2, 2.7, 3.15], 8.2)

    # 16 gaps
    doc.add_heading("16. Known gaps, legacy code and maintainability", level=1)
    add_table(doc, ["Area", "Current state / impact"], [
        ("Incomplete baseline schema", "No full initial schema dump; remote catalog is required for exact reconstruction and proof of current RLS/grants."),
        ("Inactive screens", "PayrollScreen, DocumentsScreen, ProjectsScreen and several alternate Attendance screens exist but are not in active navigation; DocumentManagementScreen is a minimal placeholder."),
        ("Historical documentation", "design_guidelines.md describes blue theme/Firebase and older implementation notes describe credential/TOTP features/files that no longer match active source. Do not treat them as production truth."),
        ("Legacy SQL", "scripts/ includes resets, direct fixes, seeding and RLS-disable utilities. These are operational hazards outside controlled recovery."),
        ("Function generations", "Multiple versions of attendance/biometric functions/views remain in migration history. PostgreSQL final state is order-dependent; never deploy a subset by filename intuition."),
        ("Edge Function config", "No supabase/config.toml is tracked; verify_jwt and local emulation settings are not reproducible from the repository."),
        ("Event media", "event-photos bucket/policy migration is missing and Event screen uses a public URL."),
        ("Messaging", "WhatsApp and general transactional SMS are intentionally disabled client placeholders; only OTP SMS is implemented server-side."),
        ("Testing", "Strong unit/regression suite but no E2E, accessibility, load, live DB policy or Edge Function CI suite."),
        ("Large modules", "Several screens/services exceed 1,000 lines, increasing regression and review risk."),
        ("Encoding", "Some comments/UI strings display mojibake in repository output (for example middle dots/emoji); enforce UTF-8 and add UI snapshot checks."),
    ], [1.6, 5.45], 8.4)

    doc.add_heading("17. Operational troubleshooting guide", level=1)
    add_table(doc, ["Symptom", "Likely cause", "Checks / response"], [
        ("OTP not received", "Inactive/unmatched phone, SMS secret/provider failure or phone throttle.", "Use generic client response; inspect Edge logs/challenge count/provider status without exposing OTP; do not bypass with static admin code."),
        ("Admin opens employee shell", "Cached profile or non-normalized/incorrect DB role.", "Refresh rpc_get_my_profile_v2/profile, normalize role, confirm active status and protected role-repair audit."),
        ("Check-in rejected", "GPS accuracy >100 m, outside distance+accuracy radius, inactive/missing office or missing policy.", "Use configuration issues RPC; verify coordinates/radius and device precise-location permission."),
        ("Attendance active after checkout", "Server did not commit checkout or policy/audit protection rejected it.", "Do not patch table directly; inspect RPC error, action request and audit records; use authorized correction/regularization."),
        ("Background tracking stale", "Permission/GPS/battery optimization/session/bridge/network issue.", "Open Tracking Setup; inspect tracking health events, active session, native foreground service and offline queue."),
        ("Field evidence stuck", "Upload/network error or metadata submit failed.", "Reconnect to retry; confirm private object and idempotency key; delete only unreferenced owner object through Edge Function."),
        ("Biometric mismatch", "Unmapped employee code, duplicate profile mapping or sync failure.", "Review unmatched list, assign unique mapping, reprocess punches, inspect sync/batch state."),
        ("Old web UI persists", "Stale Hostinger/CDN/service-worker cache.", "Verify versioned sw URL, no-cache headers for index/sw/manifest, clear only bop-hrms caches and deploy complete hashed asset set."),
        ("Deep route returns 404/HTML JS", "Apache rewrite/static asset fallback misconfiguration.", "Deploy .htaccess; static extensions must return 404, app routes must fall back to index."),
        ("Push missing", "No active token, permission denied, outbox waiting/failed or dispatcher auth/job failure.", "Check token, outbox/delivery/receipt rows, cron and Edge logs; remediate F-02 before exposing manual dispatch."),
    ], [1.6, 2.3, 3.15], 8.0)

    # Appendices
    doc.add_page_break()
    doc.add_heading("Appendix A. Complete screen inventory", level=1)
    screens = collect_screens()
    add_table(doc, ["Screen", "Route status", "Purpose", "Data / service dependencies", "Lines"], [
        (s["path"], "Active" if s["active"] else "Not active / alternate", s["description"],
         "; ".join((["tables " + ", ".join(s["tables"])] if s["tables"] else []) +
                    (["RPC " + ", ".join(s["rpcs"])] if s["rpcs"] else []) +
                    (["services " + ", ".join(s["services"])] if s["services"] else [])) or "Presentation/local only", s["lines"])
        for s in screens
    ], [1.85, 1.0, 2.35, 2.9, 0.45], 6.6)

    doc.add_heading("Appendix B. Migration inventory", level=1)
    migrations = sorted((ROOT / "supabase" / "migrations").glob("*.sql"))
    migration_rows = []
    for path in migrations:
        name = path.stem
        version, _, slug = name.partition("_")
        migration_rows.append((version, slug.replace("_", " ").capitalize(), path.name, path.stat().st_size))
    add_table(doc, ["Version", "Purpose inferred from filename", "File", "Bytes"], migration_rows, [1.2, 2.5, 3.0, 0.65], 6.8)

    doc.add_heading("Appendix C. Test inventory", level=1)
    tests = sorted([p for p in ROOT.rglob("*.test.ts") if "node_modules" not in p.parts and "dist" not in p.parts])
    test_rows = []
    for path in tests:
        text = path.read_text(encoding="utf-8", errors="replace")
        cases = len(re.findall(r"\b(?:it|test)\s*\(", text))
        test_rows.append((path.relative_to(ROOT).as_posix(), cases, text.count("\n") + 1))
    add_table(doc, ["Test file", "Declared test cases", "Lines"], test_rows, [5.7, 1.2, 0.55], 7.4)

    doc.add_heading("Appendix D. Dependency manifest", level=1)
    dep_rows = []
    for scope in ("dependencies", "devDependencies"):
        for name, version in sorted(package.get(scope, {}).items()):
            dep_rows.append((scope, name, version))
    add_table(doc, ["Scope", "Package", "Declared version/source"], dep_rows, [1.25, 3.6, 2.5], 7.4)

    doc.add_heading("Appendix E. Source manifest", level=1)
    include_dirs = ["components", "constants", "contexts", "hooks", "lib", "modules", "navigation", "screens", "services", "tasks", "types", "utils", "supabase/functions"]
    manifest = []
    for dirname in include_dirs:
        for path in sorted((ROOT / dirname).rglob("*")):
            if not path.is_file() or "__tests__" in path.parts:
                continue
            if path.suffix.lower() not in {".ts", ".tsx", ".js", ".kt", ".json"}:
                continue
            text = path.read_text(encoding="utf-8", errors="replace")
            manifest.append((path.relative_to(ROOT).as_posix(), path.suffix.lower(), text.count("\n") + 1, path.stat().st_size))
    add_table(doc, ["Source path", "Type", "Lines", "Bytes"], manifest, [5.45, 0.65, 0.55, 0.75], 6.8)

    doc.add_heading("Appendix F. Evidence and source-of-truth index", level=1)
    add_table(doc, ["Topic", "Primary source files"], [
        ("Theme", "constants/theme.ts; hooks/useTheme.ts; navigation/screenOptions.ts"),
        ("Composition/routes", "App.tsx; navigation/*.tsx"),
        ("Authentication", "contexts/AuthContext.tsx; supabase/functions/otp-auth/index.ts; utils/secureStorage.ts"),
        ("Attendance", "services/attendance.service.ts; services/canonical-attendance.service.ts; AttendanceScreen.tsx; canonical/guardrail migrations"),
        ("Field/tracking", "field-visit and current-location services; services/location/*; LocationTrackingContext; 20260720-23 and 20260803-06 migrations"),
        ("Leave", "services/leave.service.ts; 20260806020000_leave_ledger_and_comp_off.sql"),
        ("Events", "services/event.service.ts; EventManagementScreen.tsx; 20260725 and 20260810 migrations"),
        ("Biometric", "SmartOffice/import/sync services and screens; biometric Edge Functions; 20260724-20260804 migrations"),
        ("Notifications", "notification services; 20260805010000 outbox migration; notification-dispatch Edge Function"),
        ("Web/Android delivery", "App.tsx; AndroidHrmsShell.tsx; NativeShellConnector.tsx; public/sw.js; public/.htaccess; app.json/app.config.js/eas.json"),
        ("Security", "supabase/security_audit.sql; latest protected mutation/RLS migrations; npm audit snapshot; source review findings in Section 15"),
    ], [1.5, 5.55], 8.1)

    doc.add_heading("Appendix G. Glossary", level=1)
    add_table(doc, ["Term", "Meaning"], [
        ("Canonical attendance", "Single official employee/day interpretation derived from all valid evidence and policy."),
        ("Business date", "Calendar date in Asia/Kolkata, independent of client timezone."),
        ("Evidence", "Raw app, biometric, field visit, location, photo or correction fact retained separately from official result."),
        ("RPC", "PostgreSQL function exposed through Supabase; privileged mutations validate identity/role server-side."),
        ("RLS", "PostgreSQL Row Level Security policy restricting which rows a database role/user can access."),
        ("Outbox", "Database table of side effects to deliver asynchronously and idempotently."),
        ("SmartOffice", "External biometric attendance device/service integrated through server-side functions."),
        ("EAS", "Expo Application Services for builds, updates and channels."),
        ("PWA", "Installable web application with service-worker caching."),
        ("Current-location fallback", "Field visit evidence path used when a planned destination cannot be reliably selected/reached."),
    ], [1.65, 5.4], 8.4)

    doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("END OF DOCUMENT")
    r.bold = True; r.font.color.rgb = RGBColor.from_string(PURPLE); r.font.size = Pt(12)

    doc.save(OUTPUT)
    # Validate that python-docx can reopen the package and its main document part exists.
    check = Document(OUTPUT)
    if len(check.paragraphs) < 100 or len(check.tables) < 20:
        raise RuntimeError("Generated documentation failed structural validation")
    print(json.dumps({
        "output": str(OUTPUT), "bytes": OUTPUT.stat().st_size,
        "paragraphs": len(check.paragraphs), "tables": len(check.tables),
        "screens": len(screens), "services": len(services), "tables_catalogued": len(all_tables),
        "functions_catalogued": len(objects["function"]), "migrations": len(migrations), "tests": len(tests),
    }, indent=2))


if __name__ == "__main__":
    main()
