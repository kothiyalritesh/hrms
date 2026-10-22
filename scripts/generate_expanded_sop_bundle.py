from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from generate_realtyhrm_sop_doc import render_markdown, set_cell_text, set_cell_shading, add_page_number

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "realtyhrm-sop-expanded"
MD_OUT = ROOT / "docs" / "RealtyHRM_COMPLETE_SOP.md"
DOCX_OUT = ROOT / "docs" / "RealtyHRM_COMPLETE_SOP_FINAL_v5.docx"

FILES = [
    "00_GAP_ANALYSIS.md",
    "02_EMPLOYEE_MOBILE_MANUAL.md",
    "03_EMPLOYEE_WEB_MANUAL.md",
    "04_HR_PANEL_MANUAL.md",
    "05_TEAM_LEAD_MANUAL.md",
    "06_SYSTEM_ADMIN_MANUAL.md",
    "07_ATTENDANCE_BIOMETRIC_MANUAL.md",
    "08_LOCATION_FIELD_VISIT_MANUAL.md",
    "09_TECHNICAL_DATABASE_INTEGRATIONS.md",
    "10_DEPLOYMENT_TESTING_INCIDENT_RECOVERY.md",
    "17_REPORTS_ACTIVITY_SUPPORT_MANUAL.md",
    "11_PERMISSIONS_MATRIX.md",
    "12_WORKFLOW_DIAGRAMS.md",
    "13_TEST_EVIDENCE_REGISTER.md",
    "14_KNOWN_DEFECTS_REGISTER.md",
    "15_SOURCE_TRACEABILITY_MATRIX.md",
    "16_COVERAGE_MATRIX.md",
    "18_AUDIT_RECONCILIATION.md",
    "19_SCOPE_AND_APPROVAL_REGISTER.md",
    "20_SIGN_OFF_AUDIT_AND_ACTION_PLAN.md",
]

def build_markdown():
    parts = [
        "# RealtyHRM — Complete Standard Operating Procedure",
        "",
        "**Status:** Expanded repository-audited SOP. Source/deployment gaps are explicitly classified; this document does not certify production readiness.",
        "",
        "**Evidence rule:** server RPC responses, deployed database policies and audit records are authoritative. No credentials, tokens, API keys or employee records are included.",
        "",
        "## Contents",
        "",
    ]
    for index, name in enumerate(FILES, 1):
        title = name.replace(".md", "").split("_", 1)[-1].replace("_", " ").title()
        parts.append(f"{index}. {title}")
    parts.append("")
    parts.append("---")
    parts.append("")
    for name in FILES:
        parts.append((SOURCE / name).read_text(encoding="utf-8"))
        parts.append("\n---\n")
    MD_OUT.write_text("\n".join(parts), encoding="utf-8")

def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.color.rgb = RGBColor(109, 27, 109)
    styles["Heading 2"].font.color.rgb = RGBColor(109, 27, 109)
    header = section.header.paragraphs[0]
    header.text = "RealtyHRM | Complete Standard Operating Procedure"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("RealtyHRM\nComplete Standard Operating Procedure")
    run.bold = True
    run.font.size = Pt(25)
    run.font.color.rgb = RGBColor(109, 27, 109)
    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = sub.add_run("Expanded repository-audited edition")
    sr.italic = True
    sr.font.size = Pt(12)

    meta = doc.add_table(rows=3, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (label, value) in zip(meta.rows, [
        ("Source", "RealtyHRM source, migrations, Edge Functions, tests and configuration"),
        ("Date", "19 October 2026"),
        ("Readiness", "Documentation complete; production verification items remain open"),
    ]):
        set_cell_text(row.cells[0], label, bold=True, color=(255, 255, 255))
        set_cell_shading(row.cells[0], "6D1B6D")
        set_cell_text(row.cells[1], value)
    doc.add_page_break()

    for name in FILES:
        render_markdown(doc, (SOURCE / name).read_text(encoding="utf-8"))
        doc.add_page_break()
    doc.save(DOCX_OUT)

if __name__ == "__main__":
    build_markdown()
    build_docx()
    print(MD_OUT)
    print(DOCX_OUT)
