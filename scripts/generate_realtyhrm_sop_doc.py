from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
SOP_DIR = ROOT / "docs" / "realtyhrm-sop"
OUT = ROOT / "docs" / "RealtyHRM_Complete_SOP.docx"

SOP_FILES = [
    "01_EMPLOYEE_SOP.md",
    "02_HR_TEAMLEAD_SOP.md",
    "03_ADMIN_SOP.md",
    "04_ATTENDANCE_BIOMETRIC_SOP.md",
    "05_LOCATION_FIELD_VISIT_SOP.md",
    "06_LEAVE_TASKS_NOTIFICATIONS_EVENTS_SOP.md",
    "07_REPORTS_AUDIT_PRIVACY_SOP.md",
    "08_DEPLOYMENT_TROUBLESHOOTING_SOP.md",
    "09_VERIFICATION_TRACEABILITY.md",
]

def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)

def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    run = p.add_run(text.strip())
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

def add_inline(paragraph, text):
    # Keep Markdown emphasis/code readable without attempting a full Markdown renderer.
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**")
        code = part.startswith("`") and part.endswith("`")
        value = part[2:-2] if bold else part[1:-1] if code else part
        run = paragraph.add_run(value)
        run.bold = bold
        if code:
            run.font.name = "Consolas"
            run.font.size = Pt(9)

def add_table(doc, rows):
    if not rows:
        return
    header = [x.strip() for x in rows[0].strip().strip("|").split("|")]
    data = []
    for row in rows[2:]:
        if not row.strip() or set(row.replace("|", "").strip()) <= set("-:"):
            continue
        data.append([x.strip() for x in row.strip().strip("|").split("|")])
    table = doc.add_table(rows=1, cols=len(header))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for i, value in enumerate(header):
        set_cell_text(table.rows[0].cells[i], value, bold=True, color=(255, 255, 255))
        set_cell_shading(table.rows[0].cells[i], "6D1B6D")
    for values in data:
        cells = table.add_row().cells
        for i in range(len(header)):
            set_cell_text(cells[i], values[i] if i < len(values) else "")
    doc.add_paragraph()

def render_markdown(doc, text):
    lines = text.replace("\r\n", "\n").split("\n")
    i = 0
    in_code = False
    code_lines = []
    while i < len(lines):
        line = lines[i]
        if line.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.style = "No Spacing"
                run = p.add_run("\n".join(code_lines))
                run.font.name = "Consolas"
                run.font.size = Pt(8)
                run.font.color.rgb = RGBColor(70, 70, 70)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue
        if in_code:
            code_lines.append(line)
            i += 1
            continue
        if line.startswith("|") and i + 1 < len(lines) and lines[i + 1].lstrip().startswith("|"):
            rows = [line, lines[i + 1]]
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(lines[i])
                i += 1
            add_table(doc, rows)
            continue
        match = re.match(r"^(#{1,6})\s+(.*)$", line)
        if match:
            level = min(len(match.group(1)), 4)
            p = doc.add_paragraph(style=f"Heading {level}")
            add_inline(p, match.group(2).strip())
            i += 1
            continue
        bullet = re.match(r"^\s*[-*]\s+(.*)$", line)
        numbered = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if bullet or numbered:
            style = "List Bullet" if bullet else "List Number"
            p = doc.add_paragraph(style=style)
            add_inline(p, (bullet or numbered).group(1).strip())
            i += 1
            continue
        if not line.strip():
            i += 1
            continue
        p = doc.add_paragraph()
        add_inline(p, line.strip())
        i += 1

def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Page ")
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    run._r.addnext(fld)

def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    styles = doc.styles
    styles["Normal"].font.name = "Aptos"
    styles["Normal"].font.size = Pt(10)
    styles["Heading 1"].font.color.rgb = RGBColor(109, 27, 109)
    styles["Heading 1"].font.size = Pt(18)
    styles["Heading 2"].font.color.rgb = RGBColor(109, 27, 109)
    styles["Heading 2"].font.size = Pt(14)
    styles["Heading 3"].font.color.rgb = RGBColor(80, 80, 80)

    header = section.header.paragraphs[0]
    header.text = "RealtyHRM | Standard Operating Procedure"
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor(100, 100, 100)
    add_page_number(section.footer.paragraphs[0])

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.space_after = Pt(12)
    run = title.add_run("RealtyHRM\nComplete Standard Operating Procedure")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = RGBColor(109, 27, 109)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sr = subtitle.add_run("Evidence-based operational, administration, technical and verification guide")
    sr.italic = True
    sr.font.size = Pt(12)
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Table Grid"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    metadata = [
        ("System", "RealtyHRM (Expo React Native/Web + Supabase)"),
        ("Document status", "Repository-audited SOP; deployment-dependent controls are marked"),
        ("Prepared", "19 October 2026"),
        ("Source of truth", "Inspected source code, migrations, Edge Functions, tests and configuration"),
    ]
    for row, (label, value) in zip(meta.rows, metadata):
        set_cell_text(row.cells[0], label, bold=True, color=(255,255,255))
        set_cell_shading(row.cells[0], "6D1B6D")
        set_cell_text(row.cells[1], value)
    doc.add_page_break()

    doc.add_heading("Document control and evidence rules", level=1)
    p = doc.add_paragraph()
    p.add_run("Evidence classifications: ").bold = True
    p.add_run("VERIFIED means supported by inspected executable source and/or recorded test output; PARTIALLY IMPLEMENTED means a source path exists but a dependency or end-to-end confirmation is missing; DEFECTIVE means a reproducible contradiction or failure exists; UNVERIFIED means deployment or live behavior cannot be established from repository evidence alone.")
    p = doc.add_paragraph()
    p.add_run("Operational rule: ").bold = True
    p.add_run("Server RPC responses, deployed database policies and audit records are authoritative. This document does not expose credentials, tokens, API keys or employee records.")

    doc.add_heading("Contents", level=1)
    contents = [
        "1. Employee SOP", "2. HR and Team Lead SOP", "3. Admin and Permission SOP",
        "4. Attendance and Biometric SOP", "5. Location and Field-Visit SOP",
        "6. Leave, Tasks, Notifications and Events SOP", "7. Reports, Audit and Privacy SOP",
        "8. Deployment and Troubleshooting SOP", "9. Verification, Traceability and Open Questions",
    ]
    for item in contents:
        doc.add_paragraph(item, style="List Bullet")
    doc.add_page_break()

    for filename in SOP_FILES:
        path = SOP_DIR / filename
        if not path.exists():
            continue
        render_markdown(doc, path.read_text(encoding="utf-8"))
        doc.add_page_break()

    doc.add_heading("Audit release blockers", level=1)
    blockers = [
        "Full TypeScript check currently fails in scratch.tsx.",
        "Current Jest run: 45 suites passed and 7 failed; 314 tests passed and 25 failed.",
        "Production-applied migration set and deployed Edge Function versions were not live-verified.",
        "OTP validation differs between the UI and the server function.",
        "Production backup/restore evidence was not found in the inspected repository configuration.",
    ]
    for item in blockers:
        doc.add_paragraph(item, style="List Bullet")

    doc.save(OUT)
    print(OUT)

if __name__ == "__main__":
    main()
